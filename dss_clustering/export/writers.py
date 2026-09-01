"""Table and plot export writers (spec §8.11)."""

from __future__ import annotations

import csv
import io
from copy import deepcopy
from datetime import datetime, timezone
from typing import Any, Iterable

from plotly.graph_objects import Figure

from dss_clustering.styles import EXPORT_FOOTER


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def _safe_stem(stem: str) -> str:
    cleaned = "".join(c if c.isalnum() or c in "-_" else "_" for c in stem.strip())
    return cleaned.strip("_") or "export"


def _apply_plot_footer(fig: Figure) -> Figure:
    """Stamp the branded footer onto a copy of the Plotly figure."""
    stamped = deepcopy(fig)
    annotations = list(stamped.layout.annotations or ())
    annotations.append(
        {
            "text": EXPORT_FOOTER,
            "xref": "paper",
            "yref": "paper",
            "x": 0.5,
            "y": -0.08,
            "showarrow": False,
            "font": {"size": 11, "color": "#555555"},
            "xanchor": "center",
            "yanchor": "top",
        }
    )
    margin = dict(stamped.layout.margin.to_plotly_json() if stamped.layout.margin else {})
    margin["b"] = max(int(margin.get("b") or 0), 72)
    stamped.update_layout(annotations=annotations, margin=margin)
    return stamped


def _pdf_footer(canvas, doc) -> None:  # noqa: ARG001
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColorRGB(0.35, 0.35, 0.35)
    canvas.drawCentredString(doc.pagesize[0] / 2.0, 24, EXPORT_FOOTER)
    canvas.restoreState()


def write_csv(rows: Iterable[dict[str, Any]], stem: str) -> tuple[bytes, str, str]:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=["rank", "signature", "score"])
    writer.writeheader()
    for row in rows:
        writer.writerow(
            {
                "rank": row.get("rank", ""),
                "signature": row.get("signature", ""),
                "score": row.get("score", ""),
            }
        )
    data = buffer.getvalue().encode("utf-8-sig")
    return data, f"{_safe_stem(stem)}_{_stamp()}.csv", "text/csv"


def write_count_matrix_csv(
    matrix: Any,
    document_ids: Iterable[int],
    signatures: Iterable[str],
    stem: str,
    *,
    document_labels: dict[int, str] | None = None,
) -> tuple[bytes, str, str]:
    """Write the unscaled book × signature count matrix without densifying it."""
    ids = [int(document_id) for document_id in document_ids]
    columns = [str(signature) for signature in signatures]
    if tuple(matrix.shape) != (len(ids), len(columns)):
        raise ValueError("Count matrix dimensions do not match its metadata.")

    labels = document_labels or {}
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["document_id", "book", *columns])
    for row_index, document_id in enumerate(ids):
        row = matrix.getrow(row_index).toarray().ravel() if hasattr(matrix, "getrow") else matrix[row_index]
        writer.writerow(
            [
                document_id,
                labels.get(document_id, ""),
                *(int(value) for value in row),
            ]
        )
    data = buffer.getvalue().encode("utf-8-sig")
    return data, f"{_safe_stem(stem)}_{_stamp()}.csv", "text/csv"


def write_excel(rows: Iterable[dict[str, Any]], stem: str) -> tuple[bytes, str, str]:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "features"
    ws.append(["rank", "signature", "score"])
    for row in rows:
        ws.append([row.get("rank", ""), row.get("signature", ""), row.get("score", "")])
    buffer = io.BytesIO()
    wb.save(buffer)
    return (
        buffer.getvalue(),
        f"{_safe_stem(stem)}_{_stamp()}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


_EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_BLOCK_WIDTH = 3
_BLOCK_GAP = 1


def write_excel_side_by_side(
    blocks: Iterable[dict[str, Any]],
    stem: str,
) -> tuple[bytes, str, str]:
    """Write several feature tables on one sheet, in adjacent column blocks."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    packed = list(blocks)
    if not packed:
        raise ValueError("No tables to export.")

    wb = Workbook()
    ws = wb.active
    ws.title = "features"
    title_font = Font(bold=True, size=12)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDDDDD")
    center = Alignment(horizontal="center", wrap_text=True)

    for index, block in enumerate(packed):
        start = index * (_BLOCK_WIDTH + _BLOCK_GAP) + 1
        end = start + _BLOCK_WIDTH - 1
        title = str(block.get("title") or f"Table {index + 1}")
        subtitle = str(block.get("subtitle") or "")
        rows = list(block.get("rows") or [])

        ws.merge_cells(
            start_row=1, start_column=start, end_row=1, end_column=end
        )
        title_cell = ws.cell(1, start, title)
        title_cell.font = title_font
        title_cell.alignment = center

        ws.merge_cells(
            start_row=2, start_column=start, end_row=2, end_column=end
        )
        subtitle_cell = ws.cell(2, start, subtitle)
        subtitle_cell.alignment = center

        for offset, header in enumerate(("#", "Signature", "Score"), start=start):
            cell = ws.cell(3, offset, header)
            cell.font = header_font
            cell.fill = header_fill

        for row_index, row in enumerate(rows, start=4):
            ws.cell(row_index, start, row.get("rank", ""))
            ws.cell(row_index, start + 1, row.get("signature", ""))
            ws.cell(row_index, start + 2, row.get("score", ""))

        ws.column_dimensions[get_column_letter(start)].width = 6
        ws.column_dimensions[get_column_letter(start + 1)].width = 28
        ws.column_dimensions[get_column_letter(start + 2)].width = 12

    ws.freeze_panes = "A4"
    buffer = io.BytesIO()
    wb.save(buffer)
    return (
        buffer.getvalue(),
        f"{_safe_stem(stem)}_{_stamp()}.xlsx",
        _EXCEL_MIME,
    )


def write_pdf_table(
    rows: Iterable[dict[str, Any]],
    stem: str,
    *,
    title: str = "Feature table",
) -> tuple[bytes, str, str]:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        bottomMargin=48,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(title, styles["Heading2"]),
        Spacer(1, 12),
    ]
    data = [["#", "Signature", "Score"]]
    for row in rows:
        data.append(
            [
                str(row.get("rank", "")),
                str(row.get("signature", "")),
                str(row.get("score", "")),
            ]
        )
    table = Table(data, colWidths=[40, 340, 100])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return buffer.getvalue(), f"{_safe_stem(stem)}_{_stamp()}.pdf", "application/pdf"


def write_plot_image(
    fig: Figure,
    stem: str,
    fmt: str,
    *,
    width: int = 1200,
    height: int = 800,
) -> tuple[bytes, str, str]:
    fmt = fmt.lower()
    if fmt not in ("png", "pdf"):
        raise ValueError(f"Unsupported plot format: {fmt}")
    stamped = _apply_plot_footer(fig)
    data = stamped.to_image(format=fmt, width=width, height=height, scale=2)
    mime = "image/png" if fmt == "png" else "application/pdf"
    return data, f"{_safe_stem(stem)}_{_stamp()}.{fmt}", mime


def write_table(
    rows: Iterable[dict[str, Any]],
    fmt: str,
    stem: str,
    *,
    title: str = "Feature table",
) -> tuple[bytes, str, str]:
    fmt = fmt.lower()
    if fmt == "csv":
        return write_csv(rows, stem)
    if fmt == "excel":
        return write_excel(rows, stem)
    if fmt == "pdf":
        return write_pdf_table(rows, stem, title=title)
    raise ValueError(f"Unsupported table format: {fmt}")

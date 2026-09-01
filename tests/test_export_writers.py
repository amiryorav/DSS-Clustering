"""Tests for table export writers."""

from __future__ import annotations

import csv
import io

import scipy.sparse as sp
from plotly.graph_objects import Figure
from pypdf import PdfReader

from dss_clustering.export.writers import (
    write_count_matrix_csv,
    write_excel_side_by_side,
    write_table,
)
from dss_clustering.styles import EXPORT_FOOTER


ROWS = [
    {"rank": 1, "signature": "וְ", "score": 0.8},
    {"rank": 2, "signature": "הַ", "score": 0.3},
]


def test_write_csv_excel_pdf_roundtrip_bytes():
    for fmt, suffix in (("csv", ".csv"), ("excel", ".xlsx"), ("pdf", ".pdf")):
        data, filename, mime = write_table(ROWS, fmt, "demo", title="Demo")
        assert isinstance(data, (bytes, bytearray))
        assert len(data) > 10
        assert filename.endswith(suffix)
        assert "/" not in mime or mime.startswith("application/") or mime.startswith("text/")
        if fmt == "pdf":
            text = "\n".join(
                page.extract_text() or ""
                for page in PdfReader(io.BytesIO(data)).pages
            )
            assert EXPORT_FOOTER in text


def test_write_excel_side_by_side_places_tables_in_adjacent_columns():
    from openpyxl import load_workbook

    distinctive = [{"rank": 1, "signature": "א", "score": "0.9"}]
    common = [{"rank": 1, "signature": "הַ", "score": "0.4"}]
    data, filename, mime = write_excel_side_by_side(
        [
            {"title": "Top features", "subtitle": "in X", "rows": ROWS},
            {"title": "Common features", "subtitle": "in both X and Y", "rows": common},
            {"title": "Distinctive", "subtitle": "X > Y", "rows": distinctive},
            {"title": "Gap", "subtitle": "Y > X", "rows": []},
        ],
        "all four",
    )
    assert filename.endswith(".xlsx")
    assert "spreadsheet" in mime
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "Top features"
    assert ws["E1"].value == "Common features"
    assert ws["I1"].value == "Distinctive"
    assert ws["M1"].value == "Gap"
    assert ws["A2"].value == "in X"
    assert ws["E2"].value == "in both X and Y"
    assert ws["I2"].value == "X > Y"
    assert ws["A3"].value == "#"
    assert ws["B4"].value == "וְ"
    assert ws["E4"].value == 1
    assert ws["F4"].value == "הַ"
    assert ws["I4"].value == 1
    assert ws["J4"].value == "א"
    assert ws["M4"].value is None
    assert ws.freeze_panes == "A4"


def test_write_raw_count_matrix_csv():
    matrix = sp.csr_matrix([[2, 0], [1, 3]])
    data, filename, mime = write_count_matrix_csv(
        matrix,
        [11, 12],
        ["lemma=א", "lemma=ב"],
        "raw matrix",
        document_labels={11: "Book A", 12: "Book B"},
    )

    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows == [
        ["document_id", "book", "lemma=א", "lemma=ב"],
        ["11", "Book A", "2", "0"],
        ["12", "Book B", "1", "3"],
    ]
    assert filename.endswith(".csv")
    assert mime == "text/csv"


def test_write_plot_image_embeds_footer_annotation():
    fig = Figure()
    fig.add_scatter(x=[0, 1], y=[0, 1])
    from dss_clustering.export import writers as writers_mod

    stamped = writers_mod._apply_plot_footer(fig)
    texts = [str(ann["text"]) for ann in (stamped.layout.annotations or ())]
    assert EXPORT_FOOTER in texts
